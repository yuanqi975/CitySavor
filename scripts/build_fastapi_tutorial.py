"""Build a beginner-friendly FastAPI architecture tutorial as a Word document."""

from pathlib import Path
from textwrap import dedent

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
EXAMPLE = ROOT / "tutorial" / "fastapi_todo_example"
ASSETS = ROOT / "tutorial" / "fastapi_tutorial_assets"
OUTPUT = ROOT / "deliverables" / "FastAPI_架构零基础教程.docx"
CN_FONT = Path(r"C:\Windows\Fonts\msyh.ttc")
CODE_FONT = Path(r"C:\Windows\Fonts\consola.ttf")


def font(size: int, *, code: bool = False) -> ImageFont.FreeTypeFont:
    path = CODE_FONT if code and CODE_FONT.exists() else CN_FONT
    return ImageFont.truetype(str(path), size)


def write_png(path: Path, width: int, height: int, draw_content) -> None:
    image = Image.new("RGB", (width, height), "#F8FAFC")
    draw = ImageDraw.Draw(image)
    draw_content(draw, width, height)
    image.save(path)


def centered(draw: ImageDraw.ImageDraw, text: str, y: int, image_width: int, size: int, color="#0F172A") -> None:
    selected = font(size)
    box = draw.textbbox((0, 0), text, font=selected)
    draw.text(((image_width - (box[2] - box[0])) / 2, y), text, font=selected, fill=color)


def rounded_box(draw: ImageDraw.ImageDraw, rect: tuple[int, int, int, int], title: str, subtitle: str, color: str) -> None:
    x1, y1, x2, y2 = rect
    draw.rounded_rectangle(rect, radius=22, fill=color, outline="#CBD5E1", width=2)
    draw.text((x1 + 24, y1 + 18), title, font=font(30), fill="#FFFFFF")
    draw.text((x1 + 24, y1 + 64), subtitle, font=font(20), fill="#E2E8F0")


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color="#64748B") -> None:
    draw.line([start, end], fill=color, width=6)
    x, y = end
    draw.polygon([(x, y), (x - 18, y - 10), (x - 18, y + 10)], fill=color)


def render_architecture() -> None:
    def content(draw, width, _height):
        centered(draw, "FastAPI 分层架构：每一层只做自己擅长的事", 38, width, 38)
        items = [
            ("客户端", "浏览器 / Postman", "#0F766E"),
            ("路由层", "HTTP 与数据校验", "#2563EB"),
            ("服务层", "业务规则", "#7C3AED"),
            ("仓储层", "数据库查询", "#C2410C"),
            ("SQLite", "持久化存储", "#475569"),
        ]
        x = 36
        for index, (title, subtitle, color) in enumerate(items):
            rounded_box(draw, (x, 170, x + 208, 305), title, subtitle, color)
            if index < len(items) - 1:
                arrow(draw, (x + 212, 237), (x + 252, 237))
            x += 252
        draw.rounded_rectangle((175, 380, 1065, 535), radius=20, fill="#E0F2FE", outline="#7DD3FC", width=2)
        draw.text((210, 410), "核心原则：路由不直接写 SQL；服务层不关心 HTTP；仓储层不判断业务权限。", font=font(28), fill="#0C4A6E")
        draw.text((210, 460), "这样测试、改需求、换数据库时，影响范围都会更小。", font=font(25), fill="#0C4A6E")

    write_png(ASSETS / "architecture-flow.png", 1280, 600, content)


def render_request_flow() -> None:
    def content(draw, width, _height):
        centered(draw, "创建待办事项时，请求实际经历了什么？", 36, width, 38)
        rows = [
            ("1", "POST /todos", "路由读取请求 JSON，并让 Pydantic 校验 title。", "#2563EB"),
            ("2", "get_current_user", "认证依赖解析 Authorization: Bearer <JWT>。", "#7C3AED"),
            ("3", "TodoService.create", "服务层把“当前用户”和标题组合成业务动作。", "#0F766E"),
            ("4", "TodoRepository.create", "仓储层创建 ORM 对象、提交事务并返回结果。", "#C2410C"),
            ("5", "TodoRead", "FastAPI 按响应模型返回安全、稳定的 JSON。", "#475569"),
        ]
        y = 125
        for number, name, description, color in rows:
            draw.ellipse((70, y, 128, y + 58), fill=color)
            draw.text((91, y + 11), number, font=font(25), fill="#FFFFFF")
            draw.rounded_rectangle((155, y - 2, 1190, y + 62), radius=14, fill="#FFFFFF", outline="#CBD5E1", width=2)
            draw.text((180, y + 12), name, font=font(25), fill=color)
            draw.text((420, y + 14), description, font=font(20), fill="#334155")
            if number != "5":
                arrow(draw, (99, y + 62), (99, y + 88), "#94A3B8")
            y += 88

    write_png(ASSETS / "request-flow.png", 1280, 620, content)


def render_jwt_flow() -> None:
    def content(draw, width, _height):
        centered(draw, "JWT 登录与鉴权流程", 38, width, 38)
        rounded_box(draw, (65, 165, 360, 315), "登录", "邮箱 + 密码", "#2563EB")
        rounded_box(draw, (490, 165, 785, 315), "JWT", "签名后的身份凭证", "#7C3AED")
        rounded_box(draw, (915, 165, 1210, 315), "受保护接口", "读取当前用户", "#0F766E")
        arrow(draw, (370, 240), (475, 240))
        arrow(draw, (795, 240), (900, 240))
        draw.rounded_rectangle((100, 410, 1180, 525), radius=18, fill="#FEF3C7", outline="#F59E0B", width=2)
        draw.text((135, 440), "注意：JWT 不是加密的密码。它只是证明“服务器曾签发过这个用户身份”。", font=font(27), fill="#92400E")
        draw.text((135, 480), "密码只保存为哈希值；真实项目中还应把密钥放入环境变量。", font=font(23), fill="#92400E")

    write_png(ASSETS / "jwt-flow.png", 1280, 590, content)


def render_project_tree() -> None:
    tree = """fastapi_todo_example/
├── app/
│   ├── main.py                 # 应用入口：装配路由
│   ├── core/                   # 配置、数据库、安全
│   ├── models/                 # SQLAlchemy 数据表
│   ├── schemas/                # 请求/响应的数据契约
│   ├── repositories/           # 数据库读写
│   ├── services/               # 业务规则
│   ├── dependencies/           # 可复用依赖，如当前用户
│   └── routers/                # HTTP 接口
├── tests/                      # 接口行为测试
├── requirements.txt
└── Dockerfile"""

    def content(draw, width, _height):
        draw.rounded_rectangle((35, 30, width - 35, 800), radius=20, fill="#0F172A")
        draw.text((75, 65), "推荐目录结构（不是死规定，而是清晰边界）", font=font(35), fill="#E2E8F0")
        y = 140
        for line in tree.splitlines():
            draw.text((85, y), line, font=font(25, code=True), fill="#A7F3D0" if "#" not in line else "#E2E8F0")
            y += 65

    write_png(ASSETS / "project-tree.png", 1280, 850, content)


def render_code_image(source: Path, output_name: str, start: int, end: int, caption: str) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()[start - 1 : end]
    line_height = 32
    width = 1280
    height = 130 + max(1, len(lines)) * line_height + 55

    def content(draw, _width, _height):
        draw.rounded_rectangle((25, 20, width - 25, height - 20), radius=18, fill="#0F172A")
        draw.text((55, 45), caption, font=font(25), fill="#93C5FD")
        draw.text((55, 83), str(source.relative_to(EXAMPLE)).replace("\\", "/"), font=font(18, code=True), fill="#94A3B8")
        y = 130
        for index, line in enumerate(lines, start=start):
            line_font = font(19) if any("\u4e00" <= char <= "\u9fff" for char in line) else font(19, code=True)
            draw.text((60, y), f"{index:>2}", font=font(19, code=True), fill="#64748B")
            draw.text((120, y), line.expandtabs(4), font=line_font, fill="#E2E8F0")
            y += line_height

    write_png(ASSETS / output_name, width, height, content)


def create_assets() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    render_architecture()
    render_request_flow()
    render_jwt_flow()
    render_project_tree()
    render_code_image(EXAMPLE / "app" / "main.py", "code-main.png", 1, 22, "应用入口：把路由装配进 FastAPI")
    render_code_image(EXAMPLE / "app" / "routers" / "todos.py", "code-router.png", 1, 44, "路由层：把 HTTP 交给服务层")
    render_code_image(EXAMPLE / "app" / "services" / "todo_service.py", "code-service.png", 1, 47, "服务层：所有权检查与业务动作")
    render_code_image(EXAMPLE / "app" / "repositories" / "todo_repository.py", "code-repository.png", 1, 36, "仓储层：只负责数据库读写")
    render_code_image(EXAMPLE / "app" / "dependencies" / "auth.py", "code-auth.png", 1, 35, "认证依赖：把 Token 变成当前用户")
    render_code_image(EXAMPLE / "Dockerfile", "code-docker.png", 1, 7, "Dockerfile：将运行环境写成可复现的说明书")


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def add_code_text(document: Document, text: str) -> None:
    for line in text.splitlines():
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(line or " ")
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(30, 41, 59)


def add_note(document: Document, title: str, body: str, color: str = "E0F2FE") -> None:
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, color)
    heading = cell.paragraphs[0].add_run(title + "\n")
    heading.bold = True
    heading.font.color.rgb = RGBColor(12, 74, 110)
    cell.add_paragraph(body)
    document.add_paragraph()


def add_picture(document: Document, filename: str, width: float = 6.35) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(ASSETS / filename), width=Inches(width))


def setup_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    styles = document.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    styles["Normal"].font.size = Pt(10.5)
    for style_name, size, color in [("Title", 28, RGBColor(15, 23, 42)), ("Heading 1", 18, RGBColor(30, 64, 175)), ("Heading 2", 14, RGBColor(15, 118, 110))]:
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.color.rgb = color


def add_bullet(document: Document, text: str) -> None:
    document.add_paragraph(text, style="List Bullet")


def build_document() -> None:
    create_assets()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    setup_document(document)

    title = document.add_heading("FastAPI 架构零基础教程", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("用一个 Todo API 学会分层、数据库、JWT 与 Docker")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(15)
    subtitle.runs[0].font.color.rgb = RGBColor(71, 85, 105)
    document.add_paragraph()
    add_picture(document, "architecture-flow.png")
    document.add_paragraph("适合对象：从未写过 Python Web 服务，但想看懂并亲手搭建 FastAPI 项目的学习者。")
    document.add_page_break()

    document.add_heading("目录", 1)
    for chapter in [
        "1. FastAPI 是什么", "2. 先运行一个最小 API", "3. 从单文件到分层架构", "4. 数据库与 SQLAlchemy",
        "5. JWT 登录与权限控制", "6. 从请求到响应：一次完整旅程", "7. 统一错误、配置与依赖注入",
        "8. Docker 部署", "9. 测试、常见错误与下一步", "10. 完整项目代码附录",
    ]:
        document.add_paragraph(chapter)
    document.add_page_break()

    document.add_heading("FastAPI 是什么", 1)
    document.add_paragraph("FastAPI 是一个用 Python 编写 Web API 的框架。你可以把 API 理解为：前端、手机应用或其他程序向后端“提出请求”，后端再用 JSON 回答的一套约定。它的优势是：写法简洁、运行快、会自动生成交互式接口文档。")
    add_note(document, "先记住三个词", "路由：URL 对应的处理函数。Pydantic：检查请求数据是否合格的工具。依赖注入：让 FastAPI 自动提供数据库连接、当前用户等公共对象。")
    document.add_paragraph("安装并启动示例项目：")
    add_code_text(document, "cd tutorial/fastapi_todo_example\npython -m pip install -r requirements.txt\nuvicorn app.main:app --reload")
    document.add_paragraph("启动后访问 http://127.0.0.1:8000/docs，即可在浏览器中看到 FastAPI 自动生成的接口文档。")

    document.add_heading("先运行一个最小 API", 1)
    document.add_paragraph("任何 FastAPI 项目都从一个应用实例开始。下面的 /health 是“健康检查”接口：部署平台可以用它确认服务是否还活着。")
    add_picture(document, "code-main.png")
    add_note(document, "读代码的方法", "@app.get('/health') 的意思是：当收到 GET /health 请求时，执行 health 函数。函数返回的 Python 字典会自动变成 JSON。")

    document.add_heading("从单文件到分层架构", 1)
    document.add_paragraph("小练习时，把所有代码写进 main.py 没问题。但当接口变多，路由、权限、SQL 和业务规则混在一起，会让每一次修改都变得危险。分层不是为了“文件多”，而是为了让变化有边界。")
    add_picture(document, "project-tree.png")
    add_picture(document, "architecture-flow.png")
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "层"
    table.rows[0].cells[1].text = "它负责什么（不负责什么）"
    for left, right in [
        ("routers", "处理 HTTP、状态码、请求和响应；不直接写 SQL。"),
        ("services", "表达业务规则，例如“只能修改自己的待办”；不解析 URL。"),
        ("repositories", "执行数据库查询和保存；不判断谁有没有权限。"),
        ("models / schemas", "前者是数据库表，后者是 API 数据格式；不要混为一个类。"),
        ("core / dependencies", "放配置、数据库连接和可复用认证能力。"),
    ]:
        row = table.add_row().cells
        row[0].text, row[1].text = left, right

    document.add_heading("数据库与 SQLAlchemy", 1)
    document.add_paragraph("SQLite 是一个文件型数据库，非常适合学习：不用安装数据库服务，数据就保存在 todo.db 中。SQLAlchemy 则让我们用 Python 对象描述表和查询，而不是在路由里拼 SQL 字符串。")
    add_picture(document, "code-repository.png")
    add_note(document, "模型与数据模式为何分开？", "Todo（模型）包含 owner_id 等数据库细节；TodoRead（响应模式）只暴露 id、title、completed。这样即使数据库字段变化，也不会不小心把密码哈希或内部字段发给客户端。", "DCFCE7")

    document.add_heading("JWT 登录与权限控制", 1)
    document.add_paragraph("注册时，服务把密码变成不可逆的哈希值再保存。登录时比对哈希，正确才签发 JWT。之后客户端把 JWT 放在 Authorization 请求头中，接口才能知道“当前是谁”。")
    add_picture(document, "jwt-flow.png")
    add_picture(document, "code-auth.png")
    document.add_paragraph("实际项目一定不要把 secret_key 写死在代码中；应使用环境变量、密钥管理服务或部署平台的机密配置。这里写在代码中只是为了让零基础读者能一键跑通。")
    add_note(document, "401、403、404 怎么区分？", "401：没有登录或 Token 无效。403：已经登录，但明确禁止该操作。404：资源不存在。示例对“他人的待办事项”返回 404，避免泄露该资源是否存在。", "FEE2E2")

    document.add_heading("从请求到响应：一次完整旅程", 1)
    document.add_paragraph("以 POST /todos 为例。先看路由：它负责接住 HTTP 请求，但把真正的业务交给 TodoService。这样同一个服务方法将来可以被后台任务或命令行调用，而不依赖 HTTP。")
    add_picture(document, "code-router.png")
    add_picture(document, "code-service.png")
    add_picture(document, "request-flow.png")
    add_note(document, "最关键的一行", "get_by_id_and_owner(... owner_id=current_user.id) 同时按待办 ID 和当前用户查询。这比“先按 ID 找到，再忘记检查所有者”更难写错。")

    document.add_heading("统一错误、配置与依赖注入", 1)
    document.add_paragraph("依赖注入可以把重复的公共工作抽出来。示例中的 get_db 负责“打开数据库会话，接口结束后关闭”；get_current_user 负责“从 Token 得到用户”。路由函数只声明自己需要什么，FastAPI 就会自动提供。")
    add_bullet(document, "422 Unprocessable Entity：请求 JSON 不符合 Pydantic 规则，例如 title 为空。")
    add_bullet(document, "401 Unauthorized：没有提供 Bearer Token、Token 过期或签名错误。")
    add_bullet(document, "409 Conflict：注册时邮箱已经存在。")
    add_bullet(document, "404 Not Found：待办不存在，或它并不属于当前用户。")
    document.add_paragraph("配置也应集中放在 core/config.py。开发环境可以用 SQLite；生产环境切换为 PostgreSQL 时，服务层和路由层不需要改动。")

    document.add_heading("Docker 部署", 1)
    document.add_paragraph("Docker 可以把 Python 版本、依赖和启动命令写进一个镜像。这样“在我电脑上能跑”的环境也能更稳定地在服务器运行。")
    add_picture(document, "code-docker.png")
    add_code_text(document, "cd tutorial/fastapi_todo_example\ndocker build -t fastapi-todo-api .\ndocker run --rm -p 8000:8000 fastapi-todo-api")
    add_note(document, "Docker 的边界", "Docker 解决的是“如何打包和运行”，不是自动解决数据库备份、HTTPS、域名、监控或密钥管理。这些是进一步部署时需要补上的能力。", "FEF3C7")

    document.add_heading("测试、常见错误与下一步", 1)
    document.add_paragraph("测试的目标不是证明框架能工作，而是保护你的业务规则。示例测试覆盖了：服务健康、注册登录、没有 Token 不能访问、用户不能修改别人的待办事项。")
    add_code_text(document, "cd tutorial/fastapi_todo_example\npytest -v")
    document.add_heading("常见错误", 2)
    for issue, fix in [
        ("路由里直接写 SQL", "先把查询搬到 repositories，再让 services 调用它。"),
        ("把 ORM 模型直接当接口输入", "为请求和响应分别创建 Pydantic schema。"),
        ("只验证 Token，不验证资源归属", "每次读取、修改和删除时都带 owner_id 条件。"),
        ("密钥写进 Git", "使用 .env（本地）和环境变量/密钥服务（生产）。"),
        ("没有测试就重构", "先为当前行为补一条会失败的测试，再改代码。"),
    ]:
        document.add_paragraph(f"问题：{issue}。解决：{fix}")
    document.add_paragraph("下一步建议：把 SQLite 换成 PostgreSQL；使用 Alembic 管理数据库迁移；增加分页、日志、速率限制和 CI 自动测试。")

    document.add_heading("完整项目代码附录", 1)
    document.add_paragraph("以下是本教程示例项目的完整核心代码。可直接在 tutorial/fastapi_todo_example 目录中运行。为避免文档过长，__init__.py 等只起标记作用的空文件未展开。")
    for relative in [
        "app/main.py", "app/core/config.py", "app/core/database.py", "app/core/security.py",
        "app/models/base.py", "app/models/user.py", "app/models/todo.py",
        "app/schemas/user.py", "app/schemas/todo.py", "app/repositories/user_repository.py",
        "app/repositories/todo_repository.py", "app/services/auth_service.py", "app/services/todo_service.py",
        "app/dependencies/auth.py", "app/routers/auth.py", "app/routers/todos.py", "requirements.txt", "Dockerfile",
    ]:
        document.add_heading(relative, 2)
        add_code_text(document, (EXAMPLE / relative).read_text(encoding="utf-8"))

    document.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build_document()
